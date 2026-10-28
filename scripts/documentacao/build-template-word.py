import io
import json
import sys
from pathlib import Path

import fitz
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

root = Path(__file__).resolve().parents[2]
base = root / 'docs/evolucao'
pdf = fitz.open(base / 'ERS-template-original-atualizada.pdf')
proof = json.loads((base / 'evidencias/documentacao/ers-template-verificacao.json').read_text(encoding='utf-8'))
document = Document()
document.core_properties.title = 'LinkAtletas - ERS v2.5'
document.core_properties.author = 'Equipe LinkAtletas'
document.core_properties.last_modified_by = 'Lucca'
zoom = document.settings.element.find(qn('w:zoom'))
if zoom is not None:
    zoom.set(qn('w:percent'), '100')
normal = document.styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(12)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.line_spacing = Pt(12)
normal.paragraph_format.widow_control = False
backgrounds = fitz.open(base / 'ERS-template-original-atualizada.pdf')
selected = list(range(len(pdf)))
if len(sys.argv) > 1:
    selected = [int(x) - 1 for x in sys.argv[1].split(',')]


def element(name, values):
    node = OxmlElement(name)
    for key, value in values.items():
        node.set(qn('w:' + key), str(value))
    return node


def background(paragraph, data, width, height):
    run = paragraph.add_run()
    shape = run.add_picture(io.BytesIO(data), width=Pt(width), height=Pt(height))
    inline = shape._inline
    anchor = OxmlElement('wp:anchor')
    for key, value in {'distT':0,'distB':0,'distL':0,'distR':0,'simplePos':0,'relativeHeight':0,'behindDoc':1,'locked':0,'layoutInCell':1,'allowOverlap':1}.items():
        anchor.set(key, str(value))
    position = OxmlElement('wp:simplePos')
    position.set('x', '0')
    position.set('y', '0')
    anchor.append(position)
    for name in ['wp:positionH', 'wp:positionV']:
        position = OxmlElement(name)
        position.set('relativeFrom', 'page')
        offset = OxmlElement('wp:posOffset')
        offset.text = '0'
        position.append(offset)
        anchor.append(position)
    anchor.append(inline.find(qn('wp:extent')))
    anchor.append(OxmlElement('wp:wrapNone'))
    for name in ['wp:docPr', 'wp:cNvGraphicFramePr', 'a:graphic']:
        child = inline.find(qn(name))
        if child is not None:
            anchor.append(child)
    inline.getparent().replace(inline, anchor)


for index, number in enumerate(selected):
    section = document.sections[0] if index == 0 else document.add_section(WD_SECTION_START.NEW_PAGE)
    source = pdf[number]
    section.page_width = Pt(source.rect.width)
    section.page_height = Pt(source.rect.height)
    section.left_margin = Pt(0)
    section.right_margin = Pt(0)
    section.top_margin = Pt(0)
    section.bottom_margin = Pt(0)
    section.header_distance = Pt(0)
    section.footer_distance = Pt(0)
    section.header.is_linked_to_previous = False
    rectangles = [fitz.Rect(f['rect']) for f in proof['figures'] if f['page'] == number + 1]
    lines = []
    for block in source.get_text('dict')['blocks']:
        for line in block.get('lines', []):
            spans = [s for s in line['spans'] if s['text'].strip()]
            if not spans:
                continue
            middle = fitz.Rect(line['bbox']).tl + (fitz.Rect(line['bbox']).br - fitz.Rect(line['bbox']).tl) / 2
            if any(rect.contains(middle) for rect in rectangles):
                continue
            lines.append(spans)
            for span in spans:
                backgrounds[number].add_redact_annot(fitz.Rect(span['bbox']), fill=False, cross_out=False)
    backgrounds[number].apply_redactions(images=0, graphics=0)
    pixels = backgrounds[number].get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
    background(section.header.paragraphs[0], pixels.tobytes('png'), source.rect.width, source.rect.height)
    for spans in lines:
        spans.sort(key=lambda s:s['origin'][0])
        x = spans[0]['origin'][0] - .1
        size = max(s['size'] for s in spans)
        y = spans[0]['origin'][1] - size * .8
        paragraph = document.add_paragraph()
        formatting = paragraph.paragraph_format
        formatting.line_spacing = Pt(size)
        formatting.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        frame = element('w:framePr', {'w':round((source.rect.width - x) * 20),'hAnchor':'page','vAnchor':'page','x':round(x * 20),'y':round(y * 20),'wrap':'none'})
        paragraph._p.get_or_add_pPr().insert(0, frame)
        for i, span in enumerate(spans):
            if i:
                gap = span['origin'][0] - spans[i - 1]['bbox'][2]
                if gap > 2:
                    position = span['origin'][0] - x
                    formatting.tab_stops.add_tab_stop(Pt(position))
                    paragraph.add_run('\t')
            run = paragraph.add_run(span['text'].replace('\u200b', ''))
            run.font.name = 'Times New Roman' if 'TimesNewRoman' in span['font'] else 'Arial'
            run.font.size = Pt(span['size'])
            run.font.bold = 'Bold' in span['font']
            run.font.italic = 'Italic' in span['font']
            run.font.color.rgb = RGBColor.from_string(f"{span['color']:06X}")
    anchor = document.add_paragraph()
    anchor.paragraph_format.line_spacing = Pt(1)
    anchor.add_run().font.size = Pt(1)

output = base / ('ERS-completa-v2.5-editavel.docx' if len(selected) == len(pdf) else 'ERS-word-amostra.docx')
document.save(output)
print(json.dumps({'file': str(output), 'pages': len(selected)}))
